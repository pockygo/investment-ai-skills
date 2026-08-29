from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


def iter_docx_blocks(doc: DocxDocument):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def clean(value: str) -> str:
    return " ".join(value.split()).replace("|", "／")


def extract_docx(path: Path) -> list[str]:
    doc = Document(path)
    lines = []
    p_index = t_index = 0
    for block in iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(f"[P{p_index:04d}|style={block.style.name}] {text}")
            p_index += 1
            continue
        rows = [[clean(cell.text) for cell in row.cells] for row in block.rows]
        if rows:
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            lines.append(f"[TABLE {t_index}|rows={len(rows)} cols={width}]")
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * width) + " |")
            lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        t_index += 1
    for s_index, section in enumerate(doc.sections):
        for h_index, p in enumerate(section.header.paragraphs):
            if p.text.strip():
                lines.append(f"[HEADER S{s_index} P{h_index}] {p.text.strip()}")
        for f_index, p in enumerate(section.footer.paragraphs):
            if p.text.strip():
                lines.append(f"[FOOTER S{s_index} P{f_index}] {p.text.strip()}")
    return lines


def extract_pdf(path: Path) -> list[str]:
    reader = PdfReader(path)
    lines = []
    for i, page in enumerate(reader.pages, 1):
        lines.append(f"[PAGE {i}]")
        lines.append(page.extract_text() or "[NO EXTRACTABLE TEXT]")
    return lines


def extract(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8").splitlines()
    raise ValueError(f"Unsupported source type: {path}")


def main():
    parser = argparse.ArgumentParser(description="Extract diligence sources with stable page/paragraph/table markers.")
    parser.add_argument("inputs", nargs="+")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    output = []
    for raw in args.inputs:
        path = Path(raw).resolve()
        output.extend([f"# SOURCE: {path.name}", f"Path: {path}", ""])
        output.extend(extract(path))
        output.append("")
    Path(args.out).write_text("\n\n".join(output), encoding="utf-8")


if __name__ == "__main__":
    main()
