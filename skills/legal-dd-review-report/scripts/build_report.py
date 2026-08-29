from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CHINESE_DIGITS = "零一二三四五六七八九"


def chinese_number(value):
    if not 1 <= value <= 99:
        raise ValueError("Chinese parenthesized numbering supports 1-99 items")
    if value < 10:
        return CHINESE_DIGITS[value]
    tens, ones = divmod(value, 10)
    prefix = "十" if tens == 1 else CHINESE_DIGITS[tens] + "十"
    return prefix if ones == 0 else prefix + CHINESE_DIGITS[ones]


def set_run_font(run, size=None, bold=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_update_fields(doc):
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def disable_revision_tracking(doc):
    settings = doc.settings.element
    for tag in ("w:trackRevisions", "w:doNotTrackMoves", "w:doNotTrackFormatting"):
        node = settings.find(qn(tag))
        if node is not None:
            settings.remove(node)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\h \\z \\t "IPO Heading 2,1,IPO Heading 3,2" '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, end):
        run._r.append(node)


def add_paragraph(doc, text, style="IPO Normal"):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run)
    return p


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, value=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, usable_dxa=9360):
    total = sum(widths)
    if total <= 0:
        raise ValueError("Table widths must be positive")
    values = [round(usable_dxa * w / total) for w in widths]
    values[-1] += usable_dxa - sum(values)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(usable_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in values:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(values[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, block):
    headers = [str(x) for x in block.get("headers", [])]
    rows = [[str(x) for x in row] for row in block.get("rows", [])]
    if not headers:
        raise ValueError("Table block requires headers")
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("Every table row must match the header width")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10.5, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 16 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10.5)
    widths = block.get("widths") or [1] * len(headers)
    if len(widths) != len(headers):
        raise ValueError("Table widths count must match headers")
    set_table_geometry(table, widths)
    doc.add_paragraph(style="IPO Normal")


def add_blocks(doc, blocks):
    for block in blocks or []:
        if isinstance(block, str):
            add_paragraph(doc, block)
            continue
        kind = block.get("type", "paragraph")
        if kind == "paragraph":
            add_paragraph(doc, str(block.get("text", "")))
        elif kind == "list":
            numbered = bool(block.get("numbered", True))
            for i, item in enumerate(block.get("items", []), 1):
                add_paragraph(doc, (f"（{i}）" if numbered else "") + str(item))
        elif kind == "table":
            add_table(doc, block)
        else:
            raise ValueError(f"Unsupported content block: {kind}")


def validate(data):
    required = ["company_name", "department", "report_month", "diligence_cutoff", "third_party_firm", "third_party_report_title", "third_party_report_date"]
    missing = [key for key in required if not str(data.get(key, "")).strip()]
    if missing:
        raise ValueError("Missing required metadata: " + ", ".join(missing))
    if not data.get("issues"):
        raise ValueError("At least one Chapter II issue is required")
    for i, issue in enumerate(data["issues"], 1):
        for key in ("title", "description", "recommendations"):
            if not issue.get(key):
                raise ValueError(f"Issue {i} missing {key}")
        if not issue.get("risk_analysis") and (not issue.get("omit_risk_analysis") or not issue.get("omit_reason")):
            raise ValueError(f"Issue {i} requires risk_analysis or a documented omission")


def cover(doc, data):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data["company_name"])
    set_run_font(r, size=22, bold=True)
    p.paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("法务尽职调查审阅报告")
    set_run_font(r, size=18)
    for _ in range(7):
        doc.add_paragraph()
    for text in (data["department"], data["report_month"]):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=14)
    doc.add_page_break()


def background(data):
    if data.get("background_paragraphs"):
        return data["background_paragraphs"]
    short = data.get("short_name") or data["company_name"]
    scope = data.get("review_scope") or "历史沿革、主营业务、资质证照、资产等情况"
    p1 = f'{data["department"]}已对{data["company_name"]}（以下简称“公司”或“{short}”）进行了法律尽调，主要针对公司截至{data["diligence_cutoff"]}的{scope}进行了法律尽调。'
    supp = data.get("supplemental_period")
    suffix = f'，并对公司{supp}（“补充尽调期间”）的更新情况进行了补充法律尽调。' if supp else '。截至本报告出具日，尚未取得可用于更新相关事项的补充尽调材料。'
    p2 = f'{data["third_party_firm"]}已对公司截至{data["diligence_cutoff"]}的相关情况进行了法律尽调，并于{data["third_party_report_date"]}出具了《{data["third_party_report_title"]}》（“《法律尽调报告》”），我们审阅了上述《法律尽调报告》{suffix}'
    return [p1, p2]


def build(data, template, out):
    validate(data)
    shutil.copy2(template, out)
    doc = Document(out)
    clear_body(doc)
    disable_revision_tracking(doc)
    set_update_fields(doc)
    cover(doc, data)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目录")
    set_run_font(r, size=16, bold=True)
    toc = doc.add_paragraph()
    add_toc_field(toc)
    doc.add_page_break()
    add_paragraph(doc, "一、法务尽调背景情况", "IPO Heading 2")
    add_blocks(doc, background(data))
    add_paragraph(doc, "二、主要法律问题", "IPO Heading 2")
    for i, issue in enumerate(data["issues"], 1):
        add_paragraph(doc, f'（{chinese_number(i)}）{issue["title"]}', "IPO Heading 3")
        add_paragraph(doc, "1、详细描述", "IPO Heading 4")
        intro = issue.get("source_intro", "根据《法律尽调报告》：")
        if intro:
            add_paragraph(doc, intro)
        add_blocks(doc, issue["description"])
        section_no = 2
        if issue.get("risk_analysis"):
            add_paragraph(doc, f"{section_no}、风险分析", "IPO Heading 4")
            add_blocks(doc, issue["risk_analysis"])
            section_no += 1
        add_paragraph(doc, f"{section_no}、处理建议", "IPO Heading 4")
        add_blocks(doc, issue["recommendations"])
    add_paragraph(doc, "三、补充尽调期间公司情况的更新", "IPO Heading 2")
    updates = data.get("supplemental_updates") or [{"title": "补充尽调情况", "content": ["未取得补充尽调材料，相关事项截至本报告出具日的更新状态尚无法核查。"]}]
    for i, update in enumerate(updates, 1):
        add_paragraph(doc, f'（{chinese_number(i)}）{update["title"]}', "IPO Heading 3")
        add_blocks(doc, update.get("content", []))
    doc.save(out)


def main():
    parser = argparse.ArgumentParser(description="Build a legal DD review report from structured JSON.")
    parser.add_argument("input_json")
    parser.add_argument("--out", required=True)
    parser.add_argument("--template")
    args = parser.parse_args()
    script_dir = Path(__file__).resolve().parent
    template = Path(args.template).resolve() if args.template else script_dir.parent / "assets" / "review-report-template.docx"
    out = Path(args.out).resolve()
    if out == template.resolve():
        raise ValueError("Output must not overwrite the retained template")
    data = json.loads(Path(args.input_json).read_text(encoding="utf-8"))
    build(data, template, out)


if __name__ == "__main__":
    main()
