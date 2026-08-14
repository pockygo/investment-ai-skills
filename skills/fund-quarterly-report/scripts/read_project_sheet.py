"""导出项目表 xlsx 中指定公司工作表的全部非空单元格（含坐标）。

用法: python read_project_sheet.py <项目表.xlsx> [sheet名]
不传 sheet 名时列出全部工作表。
"""
import sys

import openpyxl


def main() -> None:
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(sys.argv) < 3:
        print('Sheets:', wb.sheetnames)
        return
    ws = wb[sys.argv[2]]
    for row in ws.iter_rows():
        vals = [f'{c.coordinate}={c.value!r}' for c in row if c.value is not None]
        if vals:
            print(' | '.join(vals))


if __name__ == '__main__':
    main()
