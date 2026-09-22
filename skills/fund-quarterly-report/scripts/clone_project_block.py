# -*- coding: utf-8 -*-
"""参考实现：从上一期基金报告的「项目块」生成独立项目投后报告（模式 B / 季度更新）。

设计前提（见 SKILL.md「季度更新铁律」）：
  季度更新只允许刷新「投资信息段 / (3)经营情况 / (4)退出计划及退出收益」；
  「(1)主营业务」「(2)投资时投资亮点」必须从上一期报告深拷贝原样沿用。
  因此流程是：克隆上一期报告的元素作为模板 → 只替换动态段落的文字 → 重建 body。

直接操作 document body 有三个已验证的坑，本文件已封装处理：
  1) 清空 body 时必须连末尾 sectPr 一起清掉，否则残留两个 sectPr，分节异常且页眉失效；
  2) 主报告末节的 sectPr 常不含 headerReference（靠前节继承），独立文档必须补上
     默认页眉/页脚引用，否则页眉整体丢失；
  3) 克隆段落要剥掉 bookmarkStart/End、lastRenderedPageBreak 与段落级 sectPr，
     否则出现重复书签和多余分节。

命令行：
  python3 clone_project_block.py probe  <上一期报告.docx>       # 列出 body 下标，便于选模板元素
  python3 clone_project_block.py verify <生成的文件.docx> [...] # 结构校验

作为库使用：locate / clone_para / clone_table / set_text / build / verify。
"""
import copy
import glob
import re
import sys
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
TRIANGLE_NUMID = '3'          # 三角项目符号 numId（本项目模板为 3，ilvl=0）
DEFAULT_HEADER_RID = 'rId3'   # 默认页眉（无锡国联金投和合私募基金管理有限公司）
DEFAULT_FOOTER_RID = 'rId9'   # 默认页脚（页码域）


def body_children(doc):
    return list(doc.element.body.iterchildren())


def locate(doc, text, start=0):
    """按段落文字前缀定位 body 子元素下标（模板元素索引不要写死，按文字定位更稳）。"""
    for i, ch in enumerate(body_children(doc)):
        if i < start or ch.tag != W + 'p':
            continue
        if Paragraph(ch, doc).text.strip().startswith(text):
            return i
    raise LookupError('未找到段落：' + text)


def scrub(el):
    """剥离书签、分页标记与段落级 sectPr（坑 3）。"""
    for e in list(el.iter()):
        if e.tag in (W + 'bookmarkStart', W + 'bookmarkEnd', W + 'lastRenderedPageBreak'):
            e.getparent().remove(e)
    ppr = el.find(W + 'pPr')
    if ppr is not None:
        sp = ppr.find(W + 'sectPr')
        if sp is not None:
            ppr.remove(sp)
    return el


def set_text(p, text):
    """只改第一个 run 的文字并保留其 rPr（字号/字体/加粗），删除其余 run。"""
    runs = p.findall(W + 'r')
    if not runs:
        raise ValueError('模板段落无 run，无法继承字号；请改用带 run 的参考段落')
    for r in runs[1:]:
        p.remove(r)
    r = runs[0]
    for ch in list(r):
        if ch.tag != W + 'rPr':
            r.remove(ch)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return p


def clone_para(template_el, text):
    """深拷贝模板段落并替换文字；template_el 必须是 w:p 元素。"""
    el = scrub(copy.deepcopy(template_el))
    set_text(el, text)
    return el


def clone_table(template_el, rows):
    """深拷贝模板表格并写入 rows（必须等行等列；需要增减行列时先克隆行/列再填值）。"""
    el = scrub(copy.deepcopy(template_el))
    trs = el.findall(W + 'tr')
    if len(rows) != len(trs):
        raise ValueError('行数不一致：需 %d 行，模板 %d 行' % (len(rows), len(trs)))
    for tr, row in zip(trs, rows):
        tcs = tr.findall(W + 'tc')
        if len(row) != len(tcs):
            raise ValueError('列数不一致：需 %d 列，模板 %d 列（合并单元格会减少 tc 数）'
                             % (len(row), len(tcs)))
        for tc, val in zip(tcs, row):
            ps = tc.findall(W + 'p')
            for extra in ps[1:]:
                tc.remove(extra)
            if val == '':
                continue
            set_text(ps[0], val)
    return el


def make_sectpr(tmpl_sectpr, header_rid=DEFAULT_HEADER_RID, footer_rid=DEFAULT_FOOTER_RID):
    """以模板末节 sectPr 为基础，补默认页眉/页脚引用（坑 2）。"""
    sp = copy.deepcopy(tmpl_sectpr)
    if header_rid:
        hr = OxmlElement('w:headerReference')
        hr.set(qn('w:type'), 'default')
        hr.set(qn('r:id'), header_rid)
        sp.insert(0, hr)
    if footer_rid:
        fr = OxmlElement('w:footerReference')
        fr.set(qn('w:type'), 'default')
        fr.set(qn('r:id'), footer_rid)
        sp.insert(1 if header_rid else 0, fr)
    return sp


def build(template, out_path, blocks, header_rid=DEFAULT_HEADER_RID, footer_rid=DEFAULT_FOOTER_RID):
    """用一个模板 DOCX 生成只含 blocks 的独立 DOCX。blocks 为已克隆的 w:p / w:tbl 元素列表。"""
    doc = Document(template)
    tmpl_sectpr = copy.deepcopy(body_children(doc)[-1])
    body = doc.element.body
    for ch in list(body.iterchildren()):
        body.remove(ch)                       # 坑 1：连末尾 sectPr 一起清掉
    for b in blocks:
        body.append(b)
    body.append(make_sectpr(tmpl_sectpr, header_rid, footer_rid))
    doc.save(out_path)
    return out_path


def verify(path):
    """结构校验：ZIP 完整、单一 sectPr、(1)-(4) 齐全、三角子标题、字号、意外 outlineLvl。"""
    doc = Document(path)
    print('=' * 90)
    print(path)
    ok = True
    z = zipfile.ZipFile(path)
    if z.testzip() is not None:
        ok = False
        print('  [x] ZIP 校验失败')
    sect = sum(1 for ch in body_children(doc) if ch.tag == W + 'sectPr')
    if sect != 1:
        ok = False
        print('  [x] sectPr 数量 = %d（应为 1）' % sect)
    heads, tri, outline, sizes = [], 0, [], {}
    for ch in body_children(doc):
        if ch.tag != W + 'p':
            continue
        p = Paragraph(ch, doc)
        ppr = ch.find(W + 'pPr')
        ol = ppr.find(W + 'outlineLvl') if ppr is not None else None
        npr = ppr.find(W + 'numPr') if ppr is not None else None
        numid = None
        if npr is not None:
            n = npr.find(W + 'numId')
            numid = n.get(W + 'val') if n is not None else None
        t = p.text.strip()
        if re.match(r'^\([1-4]\)', t):
            heads.append(t)
        if numid == TRIANGLE_NUMID:
            tri += 1
        if ol is not None and t:
            outline.append(t[:36])
        for r in ch.findall(W + 'r'):
            rp = r.find(W + 'rPr')
            sz = rp.find(W + 'sz') if rp is not None else None
            if sz is not None and r.text.strip():
                sizes[sz.get(W + 'val')] = sizes.get(sz.get(W + 'val'), 0) + 1
    if heads != ['(1)主营业务', '(2)投资时投资亮点', '(3)经营情况', '(4)退出计划及退出收益']:
        ok = False
        print('  [x] 小节标题异常：', heads)
    if tri != 4:
        print('  [!] 三角子标题 %d 条（模板为 4 条，按项目实际核对）' % tri)
    if len(outline) != 1:
        ok = False
        print('  [x] outlineLvl 段落应为 1 个（仅项目标题），实际：', outline)
    print('  小节：%s | 三角：%d | 字号分布(半磅)：%s | %s'
          % (heads, tri, sizes, 'PASS' if ok else 'FAIL'))
    return ok


def probe(template):
    doc = Document(template)
    for i, ch in enumerate(body_children(doc)):
        if ch.tag == W + 'p':
            p = Paragraph(ch, doc)
            ppr = ch.find(W + 'pPr')
            npr = ppr.find(W + 'numPr') if ppr is not None else None
            numid = npr.find(W + 'numId').get(W + 'val') if npr is not None and npr.find(W + 'numId') is not None else '-'
            if p.text.strip():
                print('%-4d P   numId=%-3s %s' % (i, numid, p.text.strip()[:64]))
        elif ch.tag == W + 'tbl':
            t = Table(ch, doc)
            print('%-4d TBL %dr x %dc  first=%s' % (i, len(t.rows), len(t.columns),
                                                    t.rows[0].cells[0].text[:24]))
        elif ch.tag == W + 'sectPr':
            print('%-4d SECTPR' % i)


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(0)
    cmd = sys.argv[1]
    if cmd == 'probe':
        probe(sys.argv[2])
    elif cmd == 'verify':
        files = []
        for a in sys.argv[2:]:
            files.extend(sorted(glob.glob(a)))
        sys.exit(0 if all(verify(f) for f in files) else 1)
    else:
        print(__doc__)
