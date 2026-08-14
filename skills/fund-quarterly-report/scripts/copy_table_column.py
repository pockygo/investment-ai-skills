"""在 docx 指定表格末尾深拷贝一列（保留填充色/边框/字体等全部格式），原地保存。

新列内容与被复制列相同，随后由调用方用 python-docx 改 run 文本。
列顺序按时间先后排列时：先复制旧数据列到末尾作为新期间列，再把原列改回旧期间数据。

用法: python copy_table_column.py <file.docx> <table_index> <src_col_index>
table_index、src_col_index 均从 0 开始。
"""
import sys
from copy import deepcopy

from docx import Document

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def main() -> None:
    path, t_idx, c_idx = sys.argv[1], int(sys.argv[2]), int(sys.argv[3])
    doc = Document(path)
    t = doc.tables[t_idx]

    tbl_grid = t._tbl.find(f'{W}tblGrid')
    grid_cols = tbl_grid.findall(f'{W}gridCol')
    tbl_grid.append(deepcopy(grid_cols[c_idx]))

    for row in t.rows:
        tcs = row._tr.findall(f'{W}tc')
        row._tr.append(deepcopy(tcs[c_idx]))

    doc.save(path)
    print(f'saved: table {t_idx} now has {len(doc.tables[t_idx].columns)} columns')


if __name__ == '__main__':
    main()
