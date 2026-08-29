from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


USABLE_DXA = 8300
CHINESE_DIGITS = "零一二三四五六七八九"


def chinese_number(value):
    if not 1 <= value <= 99:
        raise ValueError("Chinese numbering supports 1-99 findings")
    if value < 10:
        return CHINESE_DIGITS[value]
    tens, ones = divmod(value, 10)
    prefix = "十" if tens == 1 else CHINESE_DIGITS[tens] + "十"
    return prefix if ones == 0 else prefix + CHINESE_DIGITS[ones]


def set_font(run, size=None, bold=None, east_asia="宋体"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_document(doc):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    for tag in ("w:trackRevisions", "w:doNotTrackMoves", "w:doNotTrackFormatting"):
        node = settings.find(qn(tag))
        if node is not None:
            settings.remove(node)


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\h \\z \\t "IPO Heading 2,1,IPO Heading 3,2,IPO Heading 4,3" '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, end):
        run._r.append(node)


def add_text(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    r = p.add_run(str(text))
    set_font(r)
    return p


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        node = OxmlElement("w:cantSplit")
        node.set(qn("w:val"), "true")
        tr_pr.append(node)


def set_cell_margins(cell, value=85):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    values = [round(USABLE_DXA * float(w) / sum(widths)) for w in widths]
    values[-1] += USABLE_DXA - sum(values)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in values:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(values[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, block):
    rows = [["" if v is None else str(v) for v in row] for row in block.get("rows", [])]
    if not rows:
        raise ValueError("Table block requires rows")
    cols = len(rows[0])
    if any(len(row) != cols for row in rows):
        raise ValueError("Every table row must have equal width")
    header_rows = int(block.get("header_rows", 1))
    font_size = float(block.get("font_size", 9))
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(values):
            p = cells[ci].paragraphs[0]
            is_number = ci > 0 and len(value) < 24
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri < header_rows or is_number else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_font(r, size=font_size, bold=True if ri < header_rows else None)
        if ri < header_rows:
            set_repeat_header(table.rows[ri])
    widths = block.get("widths") or [1] * cols
    if len(widths) != cols or sum(float(x) for x in widths) <= 0:
        raise ValueError("Invalid table widths")
    set_table_geometry(table, widths)
    doc.add_paragraph(style="IPO Normal")


def add_blocks(doc, blocks):
    for block in blocks or []:
        if isinstance(block, str):
            add_text(doc, block)
            continue
        kind = block.get("type", "paragraph")
        if kind == "paragraph":
            add_text(doc, block.get("text", ""), block.get("style", "Normal"))
        elif kind == "list":
            numbered = bool(block.get("numbered", True))
            for i, item in enumerate(block.get("items", []), 1):
                prefix = f"（{i}）" if numbered else ""
                add_text(doc, prefix + str(item))
        elif kind == "table":
            add_table(doc, block)
        elif kind == "image":
            p = doc.add_paragraph(style="IPO Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(block["path"]), width=Inches(float(block.get("width_inches", 6.2))))
            if block.get("caption"):
                cp = doc.add_paragraph(style="IPO Normal")
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cp.add_run(str(block["caption"]))
                set_font(r, size=9)
        elif kind == "page_break":
            doc.add_page_break()
        else:
            raise ValueError(f"Unsupported block type: {kind}")


def heading(doc, text, level):
    return add_text(doc, text, f"IPO Heading {level}")


def require_dict(data, key, required_keys):
    value = data.get(key)
    if not isinstance(value, dict):
        raise ValueError(f"{key} must be an object")
    missing = [k for k in required_keys if k not in value]
    if missing:
        raise ValueError(f"{key} missing keys: {', '.join(missing)}")
    return value


def validate(data):
    required = ["company_name", "report_month", "diligence_cutoff", "diligence_period", "third_party_firm", "third_party_report_title", "units", "background"]
    missing = [k for k in required if not data.get(k)]
    if missing:
        raise ValueError("Missing metadata: " + ", ".join(missing))
    mm = require_dict(data, "main_matters", ["income_profit", "personnel_compensation", "cash_flow", "borrowings"])
    require_dict(mm, "income_profit", ["profit_data", "income_change"])
    require_dict(data, "financial_statements", ["balance_sheet", "income_statement", "cash_flow_statement"])
    pd = require_dict(data, "profit_details", ["sales_and_margin", "operating_expenses", "taxation"])
    require_dict(pd, "sales_and_margin", ["product_revenue", "product_margin", "trend_analysis", "major_customers"])
    require_dict(data, "balance_details", ["trading_financial_assets", "receivables", "prepayments", "inventory", "fixed_assets", "contract_liabilities", "payables"])
    if not data.get("findings"):
        raise ValueError("At least one finding is required")
    for i, finding in enumerate(data["findings"], 1):
        if not finding.get("title") or not finding.get("description") or not finding.get("response"):
            raise ValueError(f"Finding {i} requires title, description, and response")


def cover(doc, data):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(data["company_name"])
    set_font(r, size=22, bold=True)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("财务尽职调查审阅报告")
    set_font(r, size=18)
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data["report_month"])
    set_font(r, size=14)
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目录")
    set_font(r, size=16, bold=True, east_asia="黑体")
    toc = doc.add_paragraph()
    add_toc(toc)
    doc.add_page_break()


def build(data, template, out):
    validate(data)
    if Path(out).resolve() == Path(template).resolve():
        raise ValueError("Output must not overwrite the retained template")
    shutil.copy2(template, out)
    doc = Document(out)
    clear_body(doc)
    configure_document(doc)
    cover(doc, data)

    heading(doc, "一、财务尽调背景情况", 2)
    add_blocks(doc, data["background"])
    add_text(doc, f'如无特别说明，单位均为{data["units"]}。')

    mm = data["main_matters"]
    heading(doc, "二、主要尽调事项", 2)
    heading(doc, "（一）收入及利润变动情况", 3)
    heading(doc, "1、标的公司主要损益数据情况", 4)
    add_blocks(doc, mm["income_profit"]["profit_data"])
    heading(doc, "2、收入变动情况", 4)
    add_blocks(doc, mm["income_profit"]["income_change"])
    heading(doc, "（二）人员及薪酬情况", 3)
    add_blocks(doc, mm["personnel_compensation"])
    heading(doc, "（三）现金流量情况", 3)
    add_blocks(doc, mm["cash_flow"])
    heading(doc, "（四）借款合同情况", 3)
    add_blocks(doc, mm["borrowings"])

    fs = data["financial_statements"]
    heading(doc, "三、财务报表情况", 2)
    for title, key in [("（一）合并资产负债表", "balance_sheet"), ("（二）合并利润表", "income_statement"), ("（三）合并现金流量表", "cash_flow_statement")]:
        heading(doc, title, 3)
        add_blocks(doc, fs[key])

    pd = data["profit_details"]
    sm = pd["sales_and_margin"]
    heading(doc, "四、利润表明细情况", 2)
    heading(doc, "（一）销售收入及毛利率情况", 3)
    for title, key in [("1、分产品收入构成情况", "product_revenue"), ("2、分产品毛利率构成情况", "product_margin"), ("3、产品收入变动及毛利率变动分析", "trend_analysis"), ("4、主要客户销售情况", "major_customers")]:
        heading(doc, title, 4)
        add_blocks(doc, sm[key])
    heading(doc, "（二）运营费用", 3)
    add_blocks(doc, pd["operating_expenses"])
    heading(doc, "（三）税务情况", 3)
    add_blocks(doc, pd["taxation"])

    bd = data["balance_details"]
    heading(doc, "五、资产负债表明细情况", 2)
    for title, key in [("（一）交易性金融资产", "trading_financial_assets"), ("（二）应收账款", "receivables"), ("（三）预付账款", "prepayments"), ("（四）存货", "inventory"), ("（五）固定资产", "fixed_assets"), ("（六）合同负债", "contract_liabilities"), ("（七）应付账款", "payables")]:
        heading(doc, title, 3)
        add_blocks(doc, bd[key])

    heading(doc, "六、财务尽调发现的主要问题及风险揭示", 2)
    for i, finding in enumerate(data["findings"], 1):
        heading(doc, f"（{chinese_number(i)}）{finding['title']}", 3)
        add_blocks(doc, finding["description"])
        response = finding["response"]
        first, *rest = response
        if isinstance(first, str):
            add_text(doc, "应对：" + first)
        else:
            add_text(doc, "应对：")
            add_blocks(doc, [first])
        add_blocks(doc, rest)

    doc.save(out)


def main():
    parser = argparse.ArgumentParser(description="Build an internal financial DD review report")
    parser.add_argument("input_json")
    parser.add_argument("--out", required=True)
    parser.add_argument("--template")
    args = parser.parse_args()
    script_dir = Path(__file__).resolve().parent
    template = Path(args.template).resolve() if args.template else script_dir.parent / "assets" / "review-report-template.docx"
    data = json.loads(Path(args.input_json).read_text(encoding="utf-8"))
    build(data, template, Path(args.out).resolve())


if __name__ == "__main__":
    main()
