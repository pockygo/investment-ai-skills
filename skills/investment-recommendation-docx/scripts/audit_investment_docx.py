#!/usr/bin/env python3
"""Audit an investment recommendation DOCX for editability and black/white text rules."""

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--required", action="append", default=[])
    ap.add_argument("--allow-color", action="store_true")
    args = ap.parse_args()
    path = Path(args.docx)
    colors, shadings = Counter(), Counter()
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            raise SystemExit(f"Corrupt package part: {bad}")
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for name in z.namelist():
            if name == "word/document.xml" or name.startswith("word/header") or name.startswith("word/footer"):
                root = etree.fromstring(z.read(name))
                for el in root.iter(ns + "color"):
                    colors[el.get(ns + "val")] += 1
                if name == "word/document.xml":
                    for el in root.iter(ns + "shd"):
                        shadings[el.get(ns + "fill")] += 1
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        text += "\n" + "\n".join("\t".join(c.text for c in row.cells) for row in table.rows)
    forbidden_terms = ["投决报告中", "项目报告中", "项目组观点", "项目组口径", "根据PPT"]
    forbidden = {x: text.count(x) for x in forbidden_terms}
    required = {x: text.count(x) for x in args.required}
    result = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "images": len(doc.inline_shapes), "sections": len(doc.sections), "characters": len(text), "colors": dict(colors), "shadings": dict(shadings), "forbidden": forbidden, "required": required}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if any(forbidden.values()):
        raise SystemExit("Forbidden production-language phrases found")
    if any(v == 0 for v in required.values()):
        raise SystemExit("Required text missing")
    if not args.allow_color and not set(colors).issubset({"000000", "auto"}):
        raise SystemExit("Non-black Word font color found")
    if not set(shadings).issubset({"FFFFFF", "auto", None}):
        raise SystemExit("Non-white table shading found")
    print("AUDIT_OK")


if __name__ == "__main__":
    main()
