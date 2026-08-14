#!/usr/bin/env python3
"""Extract PPTX text, native tables, chart data, and picture objects to JSON/assets."""

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def walk_shapes(shapes, prefix=""):
    for idx, shape in enumerate(shapes, 1):
        key = f"{prefix}{idx}"
        yield key, shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes, key + ".")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--outdir", required=True)
    args = ap.parse_args()
    source = Path(args.pptx).resolve()
    out = Path(args.outdir).resolve()
    assets = out / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    prs = Presentation(source)
    result = {"source": str(source), "slide_width": prs.slide_width, "slide_height": prs.slide_height, "slides": []}
    for si, slide in enumerate(prs.slides, 1):
        item = {"slide": si, "texts": [], "tables": [], "charts": [], "pictures": []}
        for key, shape in walk_shapes(slide.shapes):
            base = {"key": key, "name": getattr(shape, "name", ""), "left": int(shape.left), "top": int(shape.top), "width": int(shape.width), "height": int(shape.height)}
            if getattr(shape, "has_text_frame", False) and (shape.text or "").strip():
                item["texts"].append({**base, "text": shape.text})
            if getattr(shape, "has_table", False):
                item["tables"].append({**base, "rows": [[c.text for c in r.cells] for r in shape.table.rows]})
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                rec = {**base, "chart_type": str(chart.chart_type), "categories": [], "series": []}
                try:
                    rec["categories"] = [str(x) for x in chart.plots[0].categories]
                except Exception:
                    pass
                try:
                    for series in chart.series:
                        vals = []
                        try:
                            vals = [None if v is None else float(v) for v in series.values]
                        except Exception:
                            pass
                        rec["series"].append({"name": str(series.name), "values": vals})
                except Exception:
                    pass
                item["charts"].append(rec)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    dest = assets / f"slide_{si:03d}_{key.replace('.', '_')}.{image.ext or 'bin'}"
                    dest.write_bytes(image.blob)
                    item["pictures"].append({**base, "path": str(dest)})
                except Exception as exc:
                    item["pictures"].append({**base, "error": str(exc)})
        result["slides"].append(item)
    (out / "ppt_extract.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "ppt_text.txt").write_text("\n\n".join(f"===== SLIDE {s['slide']:03d} =====\n" + "\n".join(x["text"] for x in s["texts"]) for s in result["slides"]), encoding="utf-8")
    print(json.dumps({"slides": len(result["slides"]), "texts": sum(len(s["texts"]) for s in result["slides"]), "tables": sum(len(s["tables"]) for s in result["slides"]), "charts": sum(len(s["charts"]) for s in result["slides"]), "pictures": sum(len(s["pictures"]) for s in result["slides"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
