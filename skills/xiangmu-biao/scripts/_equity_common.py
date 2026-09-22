#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""第八区（最新股权结构）公共工具：样式常量、区块定位、名册读取、表尾追加。

被 build_equity_block.py / normalize_equity_block.py 复用；也可直接 import 使用。
"""

import csv
import os
import re
from collections import OrderedDict
from copy import copy
from decimal import Decimal as D

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------- 规范样式（见 SKILL.md「样式规范」与 references/equity-structure.md） ----------------
FONT_TITLE = Font(name='微软雅黑', size=11, bold=True, color='FFFFFFFF')
FONT_HEAD = Font(name='微软雅黑', size=10, bold=True)
FONT_BODY = Font(name='微软雅黑', size=10)
FONT_NOTE = Font(name='微软雅黑', size=9, color='808080')
FONT_TOTAL = Font(name='微软雅黑', size=10, bold=True)
FILL_TITLE = PatternFill('solid', fgColor='2F5597')
FILL_HEAD = PatternFill('solid', fgColor='D9E1F2')
FILL_TOTAL = PatternFill('solid', fgColor='FCE4D6')
_SIDE = Side(style='thin', color='999999')
BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALIGN_CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_HDR = Alignment(horizontal='center', vertical='center', wrap_text=True)
FMT_PCT, FMT_AMT = '0.0000%', '#,##0.00'

BLOCK_RE = re.compile(r'^([一二三四五六七八])、')
TOTAL_RE = re.compile(r'^\s*(合\s*计|总\s*计)\s*$')
NOTE_RE = re.compile(r'^\s*(注[:：]|说明[:：]|备注[:：]|来源|数据来源|名册基准日|基准日|口径)')
HEADER_TITLE = '八、最新股权结构'


def norm_name(s):
    if not isinstance(s, str):
        return ''
    return re.sub(r'\s+', '', s.replace('（', '(').replace('）', ')')).strip()


def find_blocks(ws):
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            m = BLOCK_RE.match(v.strip())
            if m and m.group(1) not in out:
                out[m.group(1)] = r
    return out


def find_registered_capital(ws):
    """按标签在一区找「注册资本（万元）」，返回 (行号, 数值)。找不到返回 (None, None)。"""
    b = find_blocks(ws)
    for r in range(b.get('一', 4), b.get('二', ws.max_row)):
        lab = ws.cell(r, 1).value
        if isinstance(lab, str) and '注册资本' in lab:
            return r, ws.cell(r, 3).value
    return None, None


def scan_block8(ws, wsf=None):
    """解析八区。合计行 = A 列「合计/总计」或 B/C 为 =SUM() 公式。"""
    b = find_blocks(ws)
    if '八' not in b:
        return None
    b8 = b['八']
    info = {'title': b8, 'header_rows': [], 'note_rows': [], 'data': [], 'total_row': None,
            'sum_b': None, 'sum_c': None, 'stray_merges': []}

    def is_sum(v):
        return isinstance(v, str) and v.startswith('=') and 'SUM(' in v.upper()

    for r in range(b8 + 1, ws.max_row + 1):
        a, bb, c = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if a is None and bb is None and c is None:
            continue
        fb = wsf.cell(r, 2).value if wsf is not None else None
        fc = wsf.cell(r, 3).value if wsf is not None else None
        if (isinstance(a, str) and TOTAL_RE.match(a)) or is_sum(fb) or is_sum(fc):
            info['total_row'] = r
            info['sum_b'], info['sum_c'] = bb, c
            break
        if isinstance(bb, (int, float)) or isinstance(c, (int, float)):
            info['data'].append((r, a, bb, c))
        elif isinstance(a, str) and NOTE_RE.match(a):
            info['note_rows'].append(r)
        else:
            info['header_rows'].append(r)
    for m in ws.merged_cells.ranges:
        if m.min_row >= b8 and not (m.min_row == b8 and m.min_col == 1 and m.max_col == 12):
            info['stray_merges'].append(str(m))
    return info


# ---------------- 名册读取 ----------------

NAME_KEYS = ('股东', '名称', '投资人')
AMT_KEYS = ('出资', '注册资本', '股本', '持股数', '认缴', '金额', '股份数')
PCT_KEYS = ('比例', '占比', '%')


def _pick(header, keys, exclude=()):
    for i, h in enumerate(header):
        if not isinstance(h, str):
            continue
        if any(k in h for k in keys) and not any(e in h for e in exclude):
            return i
    return None


def _parse_amount(v):
    if isinstance(v, (int, float)):
        return D(str(v))
    if isinstance(v, str):
        s = v.replace(',', '').replace('，', '').strip()
        s = re.sub(r'[^\d.\-]', '', s)
        if s in ('', '-', '.'):
            return D(0)
        return D(s)
    return D(0)


def _rows_to_records(rows, name_col=None, amt_col=None, pct_col=None):
    """rows: list[list]（含表头）。自动识别列；返回 (records, header_text)。"""
    header = None
    for i, r in enumerate(rows):
        cells = [c for c in r if isinstance(c, str)]
        joined = ''.join(cells)
        if any(k in joined for k in NAME_KEYS) and any(k in joined for k in AMT_KEYS):
            header = [c if isinstance(c, str) else '' for c in r]
            body = rows[i + 1:]
            break
    if header is None:                      # 无表头：按 名称/金额[/比例] 三列假设
        header = ['股东名称', '出资额', '持股比例']
        body = rows
    nc = name_col if name_col is not None else _pick(header, NAME_KEYS)
    ac = amt_col if amt_col is not None else _pick(header, AMT_KEYS)
    pc = pct_col if pct_col is not None else _pick(header, PCT_KEYS)
    if pc == ac:
        pc = None
    if nc is None or ac is None:
        raise SystemExit('名册列识别失败，请用 --name-col/--amount-col 指定列号（从 0 开始）')

    out = []
    for r in body:
        if not r or nc >= len(r):
            continue
        name = r[nc]
        if not isinstance(name, str) or not name.strip():
            continue
        name = name.strip()
        if TOTAL_RE.match(name) or '合计' in name:
            continue
        amt = _parse_amount(r[ac] if ac < len(r) else None)
        pct = None
        if pc is not None and pc < len(r) and isinstance(r[pc], (int, float, str)):
            raw = r[pc]
            if isinstance(raw, str) and raw.strip().endswith('%'):
                try:
                    pct = float(raw.strip().rstrip('%')) / 100.0
                except ValueError:
                    pct = None
            elif isinstance(raw, (int, float)):
                pct = float(raw)
        out.append({'name': name, 'amount': amt, 'pct': pct})
    return out, ' | '.join(str(h) for h in header)


def read_register(path, sheet=None, name_col=None, amt_col=None, pct_col=None):
    """读股东名册（.xlsx/.xlsm/.csv/.docx），返回 (records, header_text)。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as f:
            rows = [list(r) for r in csv.reader(f)]
    elif ext == '.docx':
        import docx                                    # python-docx
        doc = docx.Document(path)
        if not doc.tables:
            raise SystemExit('docx 中没有表格：%s' % path)
        t = doc.tables[0]
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
    else:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    return _rows_to_records(rows, name_col, amt_col, pct_col)


def merge_shareholders(records):
    """同一股东多笔出资合并为一行；比例按合计出资/总额重算（不抄名册比例）。"""
    g = OrderedDict()
    for rec in records:
        key = norm_name(rec['name'])
        g.setdefault(key, {'name': rec['name'], 'amount': D(0)})
        g[key]['amount'] += rec['amount']
    cap = sum((v['amount'] for v in g.values()), D(0))
    if cap == 0:
        raise SystemExit('名册出资额合计为 0，无法计算比例')
    out = [{'name': v['name'], 'amount': v['amount'], 'pct': float(v['amount'] / cap)}
           for v in g.values()]
    out.sort(key=lambda x: -x['pct'])
    return out, cap


# ---------------- 表尾追加 ----------------

def style_title(ws, r, title=HEADER_TITLE):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    c = ws.cell(r, 1, title)
    c.font, c.fill, c.alignment = FONT_TITLE, FILL_TITLE, Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[r].height = 20


def _row_height(name, base=17.2, cpl=6.0):
    lines = max(1, int(len(name) / cpl + 0.999))
    return max(base, min(120.0, lines * 13.5))


def append_equity_block(ws, rows, header=('股东（发起人）', '持股比例', '认缴出资额（万元）'),
                        note=None, title=HEADER_TITLE, blank_before=True, start=None):
    """按规范把八区追加到表尾。rows: [(名称, 比例 float, 金额 Decimal|float)]。
    返回 (标题行, 表头行, 数据首行, 合计行)。"""
    if start is None:
        start = ws.max_row + (2 if blank_before else 1)
    style_title(ws, start, title)
    h = start + 1
    for i, txt in enumerate(header, 1):
        c = ws.cell(h, i, txt)
        c.font, c.fill, c.alignment, c.border = FONT_HEAD, FILL_HEAD, ALIGN_HDR, BORDER
    ws.row_dimensions[h].height = 30
    r = h
    if note:
        r += 1
        c = ws.cell(r, 1, note)
        c.font = FONT_NOTE
        ws.row_dimensions[r].height = 22
    first = r + 1
    for name, pct, amt in rows:
        r += 1
        a = ws.cell(r, 1, name)
        a.font, a.alignment, a.border = FONT_BODY, ALIGN_LEFT, BORDER
        b = ws.cell(r, 2, float(pct))
        b.font, b.alignment, b.border, b.number_format = FONT_BODY, ALIGN_CTR, BORDER, FMT_PCT
        c = ws.cell(r, 3, float(amt))
        c.font, c.alignment, c.border, c.number_format = FONT_BODY, ALIGN_CTR, BORDER, FMT_AMT
        ws.row_dimensions[r].height = _row_height(name)
    r += 1
    ws.cell(r, 1, '合计')
    ws.cell(r, 2, '=SUM(B%d:B%d)' % (first, r - 1))
    ws.cell(r, 3, '=SUM(C%d:C%d)' % (first, r - 1))
    for col in range(1, 4):
        c = ws.cell(r, col)
        c.font, c.fill, c.border = FONT_TOTAL, FILL_TOTAL, BORDER
        c.number_format = FMT_PCT if col == 2 else (FMT_AMT if col == 3 else 'General')
    ws.row_dimensions[r].height = 18
    return start, h, first, r


def backup(path):
    import datetime as _dt
    import shutil
    dst = '%s.bak-%s' % (path, _dt.datetime.now().strftime('%Y%m%d-%H%M%S'))
    shutil.copy2(path, dst)
    return dst


def ensure_writable(path):
    """回读副本常是只读权限（attachment 复制过来的文件为 0400），save 前先放开写权限。"""
    if os.path.exists(path) and not os.access(path, os.W_OK):
        os.chmod(path, 0o644)
        return True
    return False
