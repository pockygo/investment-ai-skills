#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把股东名册写成第八区「最新股权结构」并追加到指定 sheet 表尾（默认 dry-run）。

用法：
    # 先干跑看结果（不写文件）
    PYTHONIOENCODING=utf-8 python build_equity_block.py \
        --workbook 项目表.xlsx --sheet 康瑞新材 \
        --register "康瑞新材股东名册/康瑞新材-股东名册.xlsx" --base-date 2026-09-20

    # 确认无误后落盘（自动生成 .bak-时间戳 备份）
    PYTHONIOENCODING=utf-8 python build_equity_block.py ... --write

名册支持 .xlsx/.xlsm/.csv/.docx；列自动识别（股东/名称、出资/注册资本/股本/持股数/认缴、比例），
也可用 --name-col/--amount-col/--pct-col 指定列号（从 0 开始）。
八区已存在时拒绝写入，除非显式 --replace。
"""

import argparse
import sys
from decimal import Decimal as D

import openpyxl

from _equity_common import (append_equity_block, backup, ensure_writable,
                            find_registered_capital, merge_shareholders, read_register,
                            scan_block8)

HEADER_CAP = ('股东（发起人）', '持股比例', '认缴出资额（万元）')
HEADER_SHR = ('股东（发起人）', '持股比例', '持股数（万股）')


def main():
    ap = argparse.ArgumentParser(description='股东名册 → 第八区（表尾追加）')
    ap.add_argument('--workbook', required=True)
    ap.add_argument('--sheet', required=True)
    ap.add_argument('--register', required=True, help='名册文件 .xlsx/.xlsm/.csv/.docx')
    ap.add_argument('--register-sheet', default=None, help='名册为多 sheet xlsx 时指定')
    ap.add_argument('--name-col', type=int, default=None)
    ap.add_argument('--amount-col', type=int, default=None)
    ap.add_argument('--pct-col', type=int, default=None)
    ap.add_argument('--base-date', default=None, help='名册基准日 YYYY-MM-DD')
    ap.add_argument('--note', default='', help='附加注释（来源文件等）')
    ap.add_argument('--share-capital', action='store_true',
                    help='股份公司：C 列表头写「持股数（万股）」')
    ap.add_argument('--replace', action='store_true', help='八区已存在时清除后重建（谨慎）')
    ap.add_argument('--write', action='store_true', help='真正落盘（默认 dry-run）')
    ap.add_argument('--dry-run', action='store_true', help='显式干跑（默认行为，仅为可读性提供）')
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook)
    if args.sheet not in wb.sheetnames:
        sys.exit('sheet 不存在：%s（现有：%s）' % (args.sheet, ', '.join(wb.sheetnames)))
    ws = wb[args.sheet]

    records, header_text = read_register(args.register, args.register_sheet,
                                         args.name_col, args.amount_col, args.pct_col)
    if not records:
        sys.exit('名册未解析出任何股东行（表头：%s）' % header_text)
    rows, cap = merge_shareholders(records)

    reg_row, reg_val = find_registered_capital(ws)
    print('=== %s / %s ===' % (args.workbook, args.sheet))
    print('名册：%s（表头：%s）' % (args.register, header_text))
    print('股东 %d 名 · 名册合计 %s 万元 · 一区注册资本（%s）= %s'
          % (len(rows), cap, 'C%d' % reg_row if reg_row else '未找到', reg_val))
    if isinstance(reg_val, (int, float)):
        delta = cap - D(str(reg_val))
        flag = 'OK' if abs(delta) < D('0.005') else '⚠ 需核实'
        print('勾稽 ΣC vs 一区注册资本：%+.6f 万元  [%s]' % (float(delta), flag))
        if abs(delta) >= D('0.005'):
            print('  → 先判断是「登记口径已更新」还是「名册为增资前口径/漏股东」，不要直接改一区（见 references/equity-structure.md §5）')
    for i, r in enumerate(rows[:8], 1):
        print('  %2d. %-40s %8.4f%%  %s' % (i, r['name'][:40], r['pct'] * 100, r['amount']))
    if len(rows) > 8:
        print('  … 其余 %d 名略' % (len(rows) - 8))

    info = scan_block8(ws)
    if info is not None and not args.replace:
        sys.exit('该 sheet 已有八区（标题行 %d）。确认要重建请加 --replace（会先清空原八区内容再写）。' % info['title'])

    if info is not None and args.replace:
        # 先解除合并（否则 MergedCell.value 只读），再清空标题行及其下方内容
        for m in [m for m in ws.merged_cells.ranges if m.min_row >= info['title']]:
            ws.unmerge_cells(str(m))
        for r in range(info['title'], ws.max_row + 1):
            for c in range(1, 5):
                cell = ws.cell(r, c)
                if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                    cell.value = None
        start = info['title']
    else:
        start = None

    note = '注：' + ' '.join(x for x in [
        ('名册基准日：%s。' % args.base_date) if args.base_date else '',
        ('来源：%s。' % args.note) if args.note else '来源：%s。' % args.register,
        '持股比例=认缴出资额/公司注册资本，同股东多笔出资已合并；合计行取 SUM 公式。',
    ] if x)

    if not args.write:
        print('\n[DRY-RUN] 未写文件。八区将追加到第 %s 行起（标题/表头/注释/%d 名股东/合计）。'
              % (start if start else ws.max_row + 2, len(rows)))
        print('确认后加 --write 落盘（自动备份 .bak-时间戳）。')
        return 0

    b = backup(args.workbook)
    t, h, first, total = append_equity_block(
        ws, [(r['name'], r['pct'], r['amount']) for r in rows],
        header=HEADER_SHR if args.share_capital else HEADER_CAP,
        note=note, start=start)
    ensure_writable(args.workbook)
    wb.save(args.workbook)
    print('\n[WROTE] 备份：%s' % b)
    print('八区：标题 %d / 表头 %d / 明细 %d-%d / 合计 %d' % (t, h, first, total - 1, total))
    print('合计行：B%d = SUM(B%d:B%d)、C%d = SUM(C%d:C%d)'
          % (total, first, total - 1, total, first, total - 1))
    print('下一步：PYTHONIOENCODING=utf-8 python validate_workbook.py --workbook %s --sheets %s'
          % (args.workbook, args.sheet))
    return 0


if __name__ == '__main__':
    sys.exit(main())
