#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""项目表.xlsx 只读校验器（含第八区「最新股权结构」勾稽与样式合规）。

用法：
    PYTHONIOENCODING=utf-8 python validate_workbook.py --workbook 项目表.xlsx
    PYTHONIOENCODING=utf-8 python validate_workbook.py --workbook 项目表.xlsx --sheets 晋成半导体,篆芯半导体
    PYTHONIOENCODING=utf-8 python validate_workbook.py --workbook 项目表.xlsx --formulas --strict

只读，不写任何文件。--strict 时存在 ERROR 则以退出码 2 结束（供 CI / 批量任务用）。
"""

import argparse
import re
import sys
from decimal import Decimal

import openpyxl

BLOCK_NAMES = {'一': '一、项目基本信息', '二': '二、投资明细', '三': '公司融资历程',
               '四': '退出与收益', '五': '协议与风控', '六': '经营情况跟踪',
               '七': 'IRR计算', '八': '最新股权结构'}
BLOCK_RE = re.compile(r'^([一二三四五六七八])、')
TOTAL_RE = re.compile(r'^\s*(合\s*计|总\s*计)\s*$')
NOTE_RE = re.compile(r'^\s*(注[:：]|说明[:：]|备注[:：]|来源|数据来源|名册基准日|基准日|口径)')

# 我方（和合/国联系）基金关键词：用于把八区持股与②区「当前持股比例」对齐
OURS_PATTERNS = ('合创', '澄创', '星程', '合励', '守创', '诺合', '隆湫', '和合', '国联')

STD = {
    'font': '微软雅黑', 'size': 10,
    'pct_fmt': '0.0000%', 'amt_fmt': '#,##0.00',
    'title_fill': '2F5597', 'header_fill': 'D9E1F2', 'total_fill': 'FCE4D6',
    'border': 'thin',
}

problems = []   # (level, sheet, msg)


def log(level, sheet, msg):
    problems.append((level, sheet, msg))


def norm_rgb(c):
    if c is None:
        return None
    rgb = getattr(c, 'rgb', None)
    if isinstance(rgb, str) and len(rgb) >= 6:
        return rgb[-6:].upper()
    return None


def cell_fill(cell):
    f = cell.fill
    if f is None or f.fill_type is None:
        return None
    return norm_rgb(f.fgColor)


def norm_name(s):
    if not isinstance(s, str):
        return ''
    s = s.replace('（', '(').replace('）', ')')
    return re.sub(r'\s+', '', s).strip()


def base_key(name):
    """取股东/基金名的前 8 个字符作为匹配键（去掉括号内容）。"""
    n = norm_name(name)
    n = n.split('(')[0]
    return n[:8]


def find_blocks(ws):
    """返回 {区块序号字符: 标题行号}。"""
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            m = BLOCK_RE.match(v.strip())
            if m and m.group(1) not in out:
                out[m.group(1)] = r
    return out


def find_registered_capital(ws, b1, b2):
    """按标签在一区查找「注册资本」，返回 (行号, 数值)。"""
    for r in range(b1, b2):
        lab = ws.cell(r, 1).value
        if isinstance(lab, str) and '注册资本' in lab:
            return r, ws.cell(r, 3).value
    return None, None


CELL_RE = re.compile(r'^\$?([A-Z]{1,3})\$?(\d{1,5})$')
BINOP_RE = re.compile(r'^=\s*(\$?[A-Z]{1,3}\$?\d{1,5}|[\d.]+)\s*([+\-*/])\s*'
                      r'(\$?[A-Z]{1,3}\$?\d{1,5}|[\d.]+)\s*$')


def column_idx(letter):
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n


def _cell_value(ws, wsf, col, row, depth):
    if depth <= 0:
        return None
    v = ws.cell(row, column_idx(col)).value
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    f = wsf.cell(row, column_idx(col)).value if wsf is not None else None
    m = BINOP_RE.match(f) if isinstance(f, str) else None
    if not m:
        return None

    def operand(tok):
        mm = CELL_RE.match(tok)
        if not mm:
            try:
                return float(tok)
            except ValueError:
                return None
        return _cell_value(ws, wsf, mm.group(1), int(mm.group(2)), depth - 1)

    a, b = operand(m.group(1)), operand(m.group(3))
    if a is None or b is None:
        return None
    op = m.group(2)
    if op == '+':
        return a + b
    if op == '-':
        return a - b
    if op == '*':
        return a * b
    return a / b if b else None


def resolve_number(ws, wsf, row, col, depth=3):
    """取单元格数值；公式无缓存值时按简单二元公式（如 =I24/$C$16、=E24/F24）兜底计算。

    openpyxl 保存后会清空公式缓存，data_only 读出 None——此时靠兜底计算，
    否则②区「当前持股比例」这类公式列会被误判为缺失。
    """
    return _cell_value(ws, wsf, openpyxl.utils.get_column_letter(col), row, depth)


def read_block2(ws, wsf, b2, b3):
    """返回 ([(基金名, 当前持有注册资本, 当前持股比例)], 合计行号或 None)。"""
    rows, total = [], None
    for r in range(b2 + 1, b3 or ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and TOTAL_RE.match(a):
            total = r
            break
        if isinstance(a, str) and a.strip() == '基金名称':
            continue
        if isinstance(a, str) and a.strip() and not NOTE_RE.match(a):
            rows.append((a.strip(),
                         resolve_number(ws, wsf, r, 9),
                         resolve_number(ws, wsf, r, 10)))
    return rows, total


def scan_block8(ws, b8, wsf=None):
    """解析八区，返回 dict（header_rows/data/total/note_rows/sum_*）。

    合计行的判定：A 列写「合计/总计」，**或** B/C 是 `=SUM(...)` 公式
    （历史遗留：盛欣科、篆芯的合计行没有 A 列标签，只有公式）。
    必须同时传入 data_only=False 的 wsf 才能识别公式。
    """
    info = {'title': b8, 'header_rows': [], 'note_rows': [], 'data': [],
            'total_row': None, 'stray_merges': [], 'sum_b': None, 'sum_c': None}

    def is_sum_cell(v):
        return isinstance(v, str) and v.startswith('=') and 'SUM(' in v.upper()

    for r in range(b8 + 1, ws.max_row + 1):
        a, b, c = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if a is None and b is None and c is None:
            continue
        if wsf is not None:
            fb, fc = wsf.cell(r, 2).value, wsf.cell(r, 3).value
        else:
            fb = fc = None
        if (isinstance(a, str) and TOTAL_RE.match(a)) or is_sum_cell(fb) or is_sum_cell(fc):
            info['total_row'] = r
            info['sum_b'], info['sum_c'] = b, c
            break
        is_num_row = isinstance(b, (int, float)) or isinstance(c, (int, float))
        if is_num_row:
            info['data'].append((r, a, b, c))
        elif isinstance(a, str) and NOTE_RE.match(a):
            info['note_rows'].append(r)
        else:
            info['header_rows'].append(r)
    for m in ws.merged_cells.ranges:
        if m.min_row >= b8 and not (m.min_row == b8 and m.min_col == 1 and m.max_col == 12):
            info['stray_merges'].append(str(m))
    return info


def num(v):
    """紧凑显示数值：保留有效小数，不带浮点尾巴。"""
    if isinstance(v, bool) or v is None:
        return '—'
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return ('%.6f' % v).rstrip('0').rstrip('.')
    return str(v)


def pct(v):
    if isinstance(v, (int, float)):
        return '%.4f%%' % (v * 100)
    if isinstance(v, str):
        return v
    return '—'


def check_block8(ws, wsf, name, blocks, reg_row, reg_val, b2_rows):
    b8 = blocks['八']
    info = scan_block8(ws, b8, wsf)
    d = info['data']
    n = len(d)
    sum_c = sum(Decimal(str(x[3])) for x in d if isinstance(x[3], (int, float)))
    sum_b = sum(Decimal(str(x[2])) for x in d if isinstance(x[2], (int, float)))

    line = ['八@%d' % b8, '股东%d' % n, 'ΣC=%s' % num(float(sum_c)),
            'C16=%s' % num(reg_val), 'ΣB=%.4f%%' % (float(sum_b) * 100)]
    level = 'OK'

    # 1) ΣC = 一区注册资本
    if isinstance(reg_val, (int, float)):
        delta = sum_c - Decimal(str(reg_val))
        line.append('ΔC=%+.6f' % float(delta))
        if abs(delta) >= Decimal('0.005'):
            log('ERROR', name, 'Σ认缴 %s ≠ 一区注册资本 %s（差 %+.4f）'
                % (num(float(sum_c)), num(reg_val), float(delta)))
            level = 'ERROR'
        elif abs(delta) > 0:
            log('INFO', name, 'Σ认缴与一区注册资本差 %+.6f（四舍五入残留，可接受）' % float(delta))
    else:
        log('WARN', name, '一区未找到「注册资本（万元）」数值')

    # 2) ΣB = 100%
    if n:
        gap = abs(sum_b - 1)
        if gap > Decimal('0.0001'):
            log('ERROR', name, 'Σ持股比例 = %.6f%%，偏离 100%%' % (float(sum_b) * 100))
            level = 'ERROR'
        elif gap > Decimal('0.00001'):
            log('INFO', name, 'Σ持股比例 = %.6f%%（四舍五入残留）' % (float(sum_b) * 100))

    # 3) 我方基金：八区 vs ②区
    b8_names = {base_key(x[1]): x for x in d if isinstance(x[1], str)}
    matched8 = set()
    for fund, hold_c, hold_pct in b2_rows:
        if not any(p in fund for p in OURS_PATTERNS):
            continue
        key = base_key(fund)
        hit = b8_names.get(key)
        if hit is None:
            # 允许前缀更短/更长的匹配
            cands = [(k, v) for k, v in b8_names.items() if k and (k in key or key in k)]
            hit = cands[0][1] if cands else None
        if hit is None:
            log('INFO', name, '②区我方基金「%s」在八区无同名行（可能经他人管理 SPV，需人工核对穿透）' % fund)
            line.append('SPV?%s' % fund[:6])
            continue
        matched8.add(hit[0])
        if isinstance(hold_pct, (int, float)) and isinstance(hit[2], (int, float)):
            if abs(hold_pct - hit[2]) > 5e-7:
                log('WARN', name, '②区「%s」当前持股 %s vs 八区 %s —— 口径不一致，需在八区注释说明'
                    % (fund[:16], pct(hold_pct), pct(hit[2])))
                level = 'WARN' if level == 'OK' else level
            else:
                line.append('✓%s' % fund[:6])
    for r, a, b, c in d:
        if r in matched8 or not isinstance(a, str) or not any(p in a for p in OURS_PATTERNS):
            continue
        log('WARN', name, '八区有我方关联方行「%s」(%s / %s)，②区投资明细未见——确认是否需并入我方持股'
            % (a, pct(b), c))

    # 4) 合计行
    if info['total_row'] is None:
        log('ERROR', name, '八区缺少合计行')
        level = 'ERROR'
    else:
        tr = info['total_row']
        f_b = wsf.cell(tr, 2).value if wsf is not None else info['sum_b']
        f_c = wsf.cell(tr, 3).value if wsf is not None else info['sum_c']
        if not (isinstance(f_b, str) and f_b.startswith('=')) or \
           not (isinstance(f_c, str) and f_c.startswith('=')):
            log('WARN', name, '八区合计行 B/C 未使用公式（规范：=SUM()，禁止写死数值）')
            level = 'WARN' if level == 'OK' else level
        if cell_fill(ws.cell(tr, 1)) != STD['total_fill']:
            log('WARN', name, '八区合计行未使用橙底 %s' % STD['total_fill'])
            level = 'WARN' if level == 'OK' else level
        if not isinstance(ws.cell(tr, 1).value, str):
            log('WARN', name, '八区合计行 A 列缺少「合计」标签')
            level = 'WARN' if level == 'OK' else level

    # 5) 表头与样式（只检查真正承载表头文字的那一行；跨行表头的副行是 MergedCell，
    #    Excel 显示跟随左上角单元格，单独判它会产生假告警）
    hdr_main = [hr for hr in info['header_rows']
                if isinstance(ws.cell(hr, 1).value, str) and ws.cell(hr, 1).value.strip()]
    if len(info['header_rows']) > 1:
        log('WARN', name, '八区表头跨 %d 行（规范：单行表头，需人工合并）' % len(info['header_rows']))
        level = 'WARN' if level == 'OK' else level
    for hr in hdr_main:
        cell = ws.cell(hr, 1)
        if cell.font.name != STD['font'] or not cell.font.bold:
            log('WARN', name, '八区表头字体 %s/%s（规范：%s 10 加粗）'
                % (cell.font.name, cell.font.sz, STD['font']))
        if cell_fill(cell) != STD['header_fill']:
            log('WARN', name, '八区表头未使用浅蓝底 %s' % STD['header_fill'])
    for r, a, b, c in d:
        if ws.cell(r, 1).border.left.style != STD['border']:
            log('WARN', name, '八区第 %d 行边框为 %s（规范：thin 999999）' % (r, ws.cell(r, 1).border.left.style))
            break
    for r, a, b, c in d[:1]:
        if not ws.cell(r, 1).alignment.wrap_text:
            log('WARN', name, '八区 A 列未开启 wrap_text（长股东名会显示不全）')
        if isinstance(b, (int, float)) and ws.cell(r, 2).number_format not in (STD['pct_fmt'], '0.00%'):
            log('WARN', name, '八区持股比例格式 %s（规范：%s）' % (ws.cell(r, 2).number_format, STD['pct_fmt']))
        if isinstance(c, (int, float)) and ws.cell(r, 3).number_format not in (STD['amt_fmt'],):
            log('WARN', name, '八区认缴出资额格式 %r（规范：%s）' % (ws.cell(r, 3).number_format, STD['amt_fmt']))
    if info['stray_merges']:
        log('WARN', name, '八区存在多余合并区域：%s' % ', '.join(info['stray_merges']))
    if not info['note_rows']:
        log('INFO', name, '八区无注释行（建议注明名册基准日/来源/穿透口径）')

    return ' · '.join(line), level, info


def main():
    ap = argparse.ArgumentParser(description='项目表.xlsx 只读校验（含八区勾稽）')
    ap.add_argument('--workbook', required=True)
    ap.add_argument('--sheets', default='', help='逗号分隔的 sheet 名，默认全部')
    ap.add_argument('--formulas', action='store_true', help='打印全部公式单元格')
    ap.add_argument('--strict', action='store_true', help='有 ERROR 时退出码 2')
    args = ap.parse_args()

    wb_v = openpyxl.load_workbook(args.workbook, data_only=True)
    wb_f = openpyxl.load_workbook(args.workbook, data_only=False)
    wanted = [s.strip() for s in args.sheets.split(',') if s.strip()] or wb_v.sheetnames

    print('=== 校验 %s（%d sheets）===' % (args.workbook, len(wb_v.sheetnames)))
    print('\n[区块锚点]')
    print('%-12s %4s %4s %4s %4s %4s %4s %4s %4s  %7s %4s' %
          ('sheet', '一', '二', '三', '四', '五', '六', '七', '八', 'max_row', 'max_c'))
    anchors = {}
    for n in wb_v.sheetnames:
        ws = wb_v[n]
        b = find_blocks(ws)
        anchors[n] = b
        print('%-12s %4s %4s %4s %4s %4s %4s %4s %4s  %7d %4d' % (
            n, *[b.get(k, '—') for k in '一二三四五六七八'], ws.max_row, ws.max_column))

    print('\n[八区勾稽]')
    reg_info = {}
    for n in wanted:
        ws, wsf = wb_v[n], wb_f[n]
        b = anchors[n]
        if '八' not in b:
            print('%-12s 未建八区' % n)
            continue
        b1, b2 = b.get('一', 4), b.get('二', ws.max_row)
        reg_row_v, reg_val = find_registered_capital(ws, b1, b2)
        b2_rows, _ = read_block2(ws, wsf, b2, b.get('三'))
        line, level, info = check_block8(ws, wsf, n, b, reg_row_v, reg_val, b2_rows)
        reg_info[n] = (reg_row_v, reg_val)
        print('%-12s [%-5s] %s' % (n, level, line))

    if args.formulas:
        print('\n[公式单元格]')
        for n in wanted:
            forms = [('%s!%s' % (n, wsf.cell(r, c).coordinate), wsf.cell(r, c).value)
                     for r in range(1, wsf.max_row + 1) for c in range(1, wsf.max_column + 1)
                     if isinstance(wsf.cell(r, c).value, str) and wsf.cell(r, c).value.startswith('=')]
            print('  %-12s %d 个：%s' % (n, len(forms), '; '.join('%s=%s' % f for f in forms)))

    print('\n[问题清单]')
    if not problems:
        print('  无')
    for lvl, sheet, msg in problems:
        print('  %-5s %-12s %s' % (lvl, sheet, msg))

    errs = sum(1 for p in problems if p[0] == 'ERROR')
    warns = sum(1 for p in problems if p[0] == 'WARN')
    print('\n[汇总] ERROR %d · WARN %d · INFO %d' %
          (errs, warns, sum(1 for p in problems if p[0] == 'INFO')))
    if args.strict and errs:
        return 2
    return 0


if __name__ == '__main__':
    sys.exit(main())
